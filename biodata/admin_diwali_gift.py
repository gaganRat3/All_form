import io
import openpyxl
from openpyxl.utils import get_column_letter
from django.contrib import admin
from django.http import HttpResponse

from .models_diwali_gift import DiwaliGiftRegistration


@admin.register(DiwaliGiftRegistration)
class DiwaliGiftRegistrationAdmin(admin.ModelAdmin):
    list_display = (
        'candidate_name',
        'dob',
        'current_city',
        'candidate_mobile',
        'parents_mobile',
        'contact_person',
        'contact_mobile',
        'pincode',
        'submitted_at',
    )
    search_fields = (
        'candidate_name',
        'current_city',
        'candidate_mobile',
        'parents_mobile',
        'contact_person',
        'contact_mobile',
        'pincode',
        'address',
    )
    list_filter = ('current_city', 'submitted_at')
    readonly_fields = ('submitted_at',)
    ordering = ['-submitted_at']
    actions = ['export_selected_to_excel']

    @admin.action(description="Export Selected to Excel")
    def export_selected_to_excel(self, request, queryset):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Diwali Gift Registrations"

        headers = [
            "ID",
            "Candidate Name",
            "Date of Birth",
            "Current City",
            "Candidate Mobile",
            "Parents Mobile",
            "Contact Person",
            "Delivery Address",
            "Courier Mobile",
            "Pincode",
            "Submitted At",
        ]
        ws.append(headers)

        for obj in queryset:
            ws.append([
                obj.id,
                obj.candidate_name,
                obj.dob,
                obj.current_city,
                obj.candidate_mobile,
                obj.parents_mobile,
                obj.contact_person,
                obj.address,
                obj.contact_mobile,
                obj.pincode,
                obj.submitted_at.strftime('%Y-%m-%d %H:%M:%S') if obj.submitted_at else '',
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response['Content-Disposition'] = 'attachment; filename=Diwali_Gift_Registrations.xlsx'
        return response

